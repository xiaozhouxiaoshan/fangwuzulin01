# -*- coding: utf-8 -*-
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
import subprocess

from docx import Document
from docx.shared import Pt


def normalize(text: str) -> str:
    return (text or "").strip()


def find_para_index_exact(doc: Document, text: str, start: int = 0) -> int:
    target = text.strip()
    for i in range(start, len(doc.paragraphs)):
        if normalize(doc.paragraphs[i].text) == target:
            return i
    return -1


def find_para_index_contains(doc: Document, needle: str, start: int = 0) -> int:
    for i in range(start, len(doc.paragraphs)):
        if needle in (doc.paragraphs[i].text or ""):
            return i
    return -1


def delete_paragraph(paragraph) -> None:
    el = paragraph._element
    el.getparent().remove(el)


def delete_table(table) -> None:
    el = table._element
    el.getparent().remove(el)


@dataclass(frozen=True)
class RunStyle:
    font_size_pt: float | None = None
    bold: bool | None = None


def capture_run_style(paragraph) -> RunStyle:
    run = paragraph.runs[0] if paragraph.runs else None
    size = run.font.size.pt if run and run.font.size else None
    bold = run.bold if run else None
    return RunStyle(font_size_pt=size, bold=bold)


def apply_run_style(paragraph, style: RunStyle) -> None:
    if not paragraph.runs:
        return
    run = paragraph.runs[0]
    if style.font_size_pt is not None:
        run.font.size = Pt(style.font_size_pt)
    if style.bold is not None:
        run.bold = style.bold


def insert_para_before(anchor_para, text: str, style: RunStyle | None) -> None:
    p = anchor_para.insert_paragraph_before(text)
    if style is not None:
        apply_run_style(p, style)


def insert_table_before(doc: Document, anchor_para, rows: int, cols: int, style_name: str = "Table Grid"):
    table = doc.add_table(rows=rows, cols=cols)
    table.style = style_name
    anchor_para._element.addprevious(table._element)
    return table


def extract_versions() -> dict[str, str]:
    versions: dict[str, str] = {}

    try:
        p = subprocess.run(["java", "-version"], capture_output=True, text=True, check=True)
        out = (p.stdout or "") + (p.stderr or "")
        line = next((l.strip() for l in out.splitlines() if "version" in l), "")
        if line:
            versions["java_runtime"] = line
    except Exception:
        pass

    try:
        p = subprocess.run(["mvn", "-v"], capture_output=True, text=True, check=True)
        out = (p.stdout or "") + (p.stderr or "")
        if out.splitlines():
            versions["maven"] = out.splitlines()[0].strip()
    except Exception:
        pass

    try:
        p = subprocess.run(["node", "-v"], capture_output=True, text=True, check=True)
        versions["node"] = (p.stdout or "").strip()
    except Exception:
        pass

    try:
        p = subprocess.run(["npm", "-v"], capture_output=True, text=True, check=True)
        versions["npm"] = (p.stdout or "").strip()
    except Exception:
        pass

    return versions


def parse_pom_versions(pom_path: Path) -> dict[str, str]:
    text = pom_path.read_text(encoding="utf-8", errors="ignore")
    out: dict[str, str] = {}

    parent_ver = re.search(r"<parent>.*?<version>([^<]+)</version>.*?</parent>", text, flags=re.S)
    if parent_ver:
        out["spring_boot_parent"] = parent_ver.group(1).strip()

    java_ver = re.search(r"<java\\.version>([^<]+)</java\\.version>", text)
    if java_ver:
        out["java_target"] = java_ver.group(1).strip()

    mysql_conn = re.search(
        r"<groupId>mysql</groupId>\s*<artifactId>mysql-connector-java</artifactId>\s*<version>([^<]+)</version>",
        text,
        flags=re.S,
    )
    if mysql_conn:
        out["mysql_connector_java"] = mysql_conn.group(1).strip()

    return out


def parse_package_versions(pkg_json: Path) -> dict[str, str]:
    text = pkg_json.read_text(encoding="utf-8", errors="ignore")

    def find(dep: str) -> str | None:
        m = re.search(rf"\"{re.escape(dep)}\"\s*:\s*\"([^\"]+)\"", text)
        return m.group(1) if m else None

    out: dict[str, str] = {}
    for dep in ["vue", "element-ui", "echarts", "axios", "@vue/cli-service"]:
        v = find(dep)
        if v:
            out[dep] = v
    return out


def parse_sql_schema(sql_path: Path, table_names: list[str]) -> dict[str, list[dict[str, str]]]:
    sql = sql_path.read_text(encoding="utf-8", errors="ignore")
    schema: dict[str, list[dict[str, str]]] = {}

    for table in table_names:
        start_pat = f"CREATE TABLE IF NOT EXISTS `{table}`"
        start = sql.find(start_pat)
        if start == -1:
            continue
        open_idx = sql.find("(", start)
        if open_idx == -1:
            continue

        # Extract the table definition body by balancing parentheses; ignore parentheses inside strings.
        i = open_idx + 1
        depth = 1
        in_str = False
        esc = False
        while i < len(sql) and depth > 0:
            ch = sql[i]
            if in_str:
                if esc:
                    esc = False
                elif ch == "\\":
                    esc = True
                elif ch == "'":
                    in_str = False
            else:
                if ch == "'":
                    in_str = True
                elif ch == "(":
                    depth += 1
                elif ch == ")":
                    depth -= 1
            i += 1

        if depth != 0:
            continue

        body = sql[open_idx + 1 : i - 1]
        cols: list[dict[str, str]] = []
        for line in body.splitlines():
            line = line.strip().rstrip(",")
            if not line.startswith("`"):
                continue
            cm = re.match(r"`([^`]+)`\s+([^\s]+)(.*)", line)
            if not cm:
                continue
            name = cm.group(1)
            typ = cm.group(2)
            rest = cm.group(3)
            not_null = "NOT NULL" in rest

            length = ""
            lm = re.search(r"\((\d+)\)", typ)
            if lm:
                length = lm.group(1)
                base_type = re.sub(r"\(\d+\)", "", typ)
            else:
                base_type = typ

            default = ""
            dm = re.search(r"DEFAULT\s+([^\s]+)", rest)
            if dm:
                default = dm.group(1).strip()

            comment = ""
            com = re.search(r"COMMENT\s+'([^']*)'", rest)
            if com:
                comment = com.group(1).strip()

            cols.append(
                {
                    "name": name,
                    "type": base_type,
                    "length": length,
                    "nullable": "否" if not_null else "是",
                    "default": default,
                    "comment": comment,
                }
            )

        schema[table] = cols

    return schema


def count_sql_inserts(sql_path: Path) -> dict[str, int]:
    sql = sql_path.read_text(encoding="utf-8", errors="ignore")
    pat = re.compile(r"INSERT INTO `(?P<table>[^`]+)`[^;]*?VALUES\s*(?P<values>\(.*?\));", re.S)
    counts: dict[str, int] = {}

    for m in pat.finditer(sql):
        table = m.group("table")
        values = m.group("values")
        level = 0
        in_str = False
        esc = False
        tuples = 0
        for ch in values:
            if in_str:
                if esc:
                    esc = False
                elif ch == "\\":
                    esc = True
                elif ch == "'":
                    in_str = False
                continue
            if ch == "'":
                in_str = True
                continue
            if ch == "(":
                if level == 0:
                    tuples += 1
                level += 1
            elif ch == ")":
                level = max(0, level - 1)
        counts[table] = counts.get(table, 0) + tuples

    return counts


def set_cell_text(cell, text: str) -> None:
    cell.text = ""
    cell.paragraphs[0].add_run(text)


def update_cover_table(doc: Document, title: str) -> None:
    if not doc.tables:
        return
    t = doc.tables[0]

    # row 0: repeated title cells
    for c in range(1, len(t.rows[0].cells)):
        set_cell_text(t.rows[0].cells[c], title)

    # Avoid fake personal info: keep structure but set placeholders
    def safe_set(row: int, col: int, value: str) -> None:
        if row < len(t.rows) and col < len(t.rows[row].cells):
            set_cell_text(t.rows[row].cells[col], value)

    safe_set(1, 1, "（待填写）")  # 二级学院
    safe_set(2, 1, "（待填写）")  # 姓名
    safe_set(2, 3, "（待填写）")  # 学号
    safe_set(3, 1, "（待填写）")  # 专业
    safe_set(3, 3, "（待填写）")  # 届别
    safe_set(4, 1, "（待填写）")  # 校内指导教师
    safe_set(4, 3, "（待填写）")  # 职称
    safe_set(5, 1, "（待填写）")  # 校外指导教师
    safe_set(5, 3, "（待填写）")  # 职称


def main() -> None:
    work_dir = Path(r"D:\BaiduNetdiskDownload\T020源码\T020源码\_doc_work")
    template_path = Path(r"C:\Users\15964\Desktop\廖富豪-新生军训管理系统的设计与实现（最终稿）.docx")
    sql_path = Path(r"D:\BaiduNetdiskDownload\T020源码\T020源码\zhihuifangwuzulin.sql")
    pom_path = Path(r"D:\BaiduNetdiskDownload\T020源码\T020源码\back\pom.xml")
    pkg_path = Path(r"D:\BaiduNetdiskDownload\T020源码\T020源码\front\package.json")

    out_path = Path(r"C:\Users\15964\Desktop\基于Spring Boot与Vue的智慧房屋租赁管理系统设计与实现_模板版.docx")
    work_copy = work_dir / "thesis_template_based.docx"

    title = "智慧房屋租赁管理系统的设计与实现"
    doc = Document(str(template_path))

    # Remove all non-cover tables (template contains old system schemas).
    for t in list(doc.tables)[1:]:
        delete_table(t)

    update_cover_table(doc, title=title)

    # Capture prototype styles BEFORE deleting the original body
    idx_h1 = find_para_index_exact(doc, "1. 引言", start=100)
    idx_h2 = find_para_index_exact(doc, "1.1 研究背景", start=100)
    idx_normal = find_para_index_contains(doc, "近年来，随着高校扩招政策", start=100)
    h1_style = capture_run_style(doc.paragraphs[idx_h1]) if idx_h1 != -1 else RunStyle(font_size_pt=16.0, bold=None)
    h2_style = capture_run_style(doc.paragraphs[idx_h2]) if idx_h2 != -1 else RunStyle(font_size_pt=15.0, bold=True)
    normal_style = capture_run_style(doc.paragraphs[idx_normal]) if idx_normal != -1 else RunStyle(font_size_pt=14.0, bold=None)

    # Replace abstracts (be conservative; no unverified claims)
    cn_abs = (
        "随着房屋租赁需求的不断增长，传统依赖人工登记和线下沟通的管理方式在信息维护、流程跟踪和数据统计等方面效率较低，"
        "难以满足租赁业务的精细化管理需求。为提升房屋租赁业务的信息化管理水平，本文设计并实现了一套基于 B/S 架构的智慧房屋租赁管理系统。"
        "系统后端采用 Spring Boot 构建 REST 接口，结合 MyBatis-Plus 与 MySQL 完成数据持久化；前端基于 Vue 与 Element UI 实现交互页面，"
        "并引入 ECharts 完成统计数据可视化。系统围绕管理员、房主与租客三类角色，提供房屋信息管理、预约看房、合同信息、报修与维修处理、"
        "房屋评价、留言与公告等功能，形成较完整的租赁业务闭环。本文同时给出数据库脚本中的初始化数据统计与关键表结构设计，为系统复现与验收提供依据。"
    )
    cn_kw = "关键词：房屋租赁；B/S架构；Spring Boot；Vue；MySQL；ECharts"
    en_abs = (
        "With the growing demand for housing rentals, traditional manual and offline management methods are inefficient in information maintenance, "
        "workflow tracking, and data analysis. This thesis designs and implements a smart housing rental management system based on the B/S architecture. "
        "The backend is built with Spring Boot and uses MyBatis-Plus with MySQL for persistence. The frontend is implemented with Vue and Element UI, "
        "and ECharts is integrated for statistical visualization. The system serves three roles (administrator, landlord, and tenant) and supports house listing management, "
        "viewing appointments, contract management, repairs and maintenance, evaluations, messages, and announcements, forming a complete rental workflow. "
        "This thesis also provides reproducible initial dataset statistics and key table schemas based on the database script."
    )
    en_kw = "Keywords: Housing Rental; B/S Architecture; Spring Boot; Vue; MySQL; ECharts"

    idx_cn = find_para_index_exact(doc, "摘  要")
    if idx_cn != -1 and idx_cn + 2 < len(doc.paragraphs):
        doc.paragraphs[idx_cn + 1].text = cn_abs
        doc.paragraphs[idx_cn + 2].text = cn_kw
        apply_run_style(doc.paragraphs[idx_cn + 1], normal_style)
        apply_run_style(doc.paragraphs[idx_cn + 2], normal_style)

    idx_en = find_para_index_exact(doc, "Abstract")
    if idx_en != -1 and idx_en + 2 < len(doc.paragraphs):
        doc.paragraphs[idx_en + 1].text = en_abs
        doc.paragraphs[idx_en + 2].text = en_kw
        apply_run_style(doc.paragraphs[idx_en + 1], normal_style)
        apply_run_style(doc.paragraphs[idx_en + 2], normal_style)

    # Prepare real data
    runtime_versions = extract_versions()
    pom_versions = parse_pom_versions(pom_path) if pom_path.exists() else {}
    pkg_versions = parse_package_versions(pkg_path) if pkg_path.exists() else {}
    insert_counts = count_sql_inserts(sql_path) if sql_path.exists() else {}

    key_tables = [
        ("fangwuxinxi", "房屋信息"),
        ("yuyuekanfang", "预约看房"),
        ("hetongxinxi", "合同信息"),
        ("fangwubaoxiu", "房屋报修"),
        ("weixiuchuli", "维修处理"),
        ("fangwupingjia", "房屋评价"),
        ("woyaodangfangzhu", "我要当房主"),
        ("fangzhu", "房主"),
        ("yonghu", "租客用户"),
        ("news", "公告"),
        ("messages", "留言板"),
        ("storeup", "收藏"),
    ]
    schema_tables = [
        "fangwuxinxi",
        "yuyuekanfang",
        "hetongxinxi",
        "fangwubaoxiu",
        "weixiuchuli",
        "woyaodangfangzhu",
    ]
    schema = parse_sql_schema(sql_path, schema_tables) if sql_path.exists() else {}

    # Delete original body content (from "1. 引言" body occurrence to before "参考文献")
    idx_body = find_para_index_exact(doc, "1. 引言", start=100)
    idx_ref = find_para_index_exact(doc, "参考文献", start=0)
    if idx_body != -1 and idx_ref != -1 and idx_ref > idx_body:
        for i in range(idx_ref - 1, idx_body - 1, -1):
            delete_paragraph(doc.paragraphs[i])

    # Anchor for inserting new body
    idx_ref = find_para_index_exact(doc, "参考文献", start=0)
    ref_para = doc.paragraphs[idx_ref] if idx_ref != -1 else doc.paragraphs[-1]

    def h1(text: str) -> None:
        insert_para_before(ref_para, text, h1_style)

    def h2(text: str) -> None:
        insert_para_before(ref_para, text, h2_style)

    def p(text: str) -> None:
        insert_para_before(ref_para, text, normal_style)

    # Chapter 1
    h1("1. 引言")
    h2("1.1 研究背景")
    p(
        "房屋租赁看似是“找房—签约—入住”三步，但落到管理细节时，至少要覆盖房源信息维护、预约看房、合同条款与租期、"
        "押金与租金状态、租后报修处理、评价与纠纷留痕等环节。对租客而言，这些环节分散在聊天记录、纸质合同和电话沟通里；"
        "对房主与管理方而言，最麻烦的不是数据量，而是流程不可追溯：同一套房源在不同时间被不同人咨询、预约、签约与退租，"
        "如果没有统一的信息载体，后续很难对账和回溯。"
    )
    p(
        "随着“线上预约 + 线下带看”的常态化，租赁业务的时间节点越来越密集，且信息更新频繁（例如房屋状态从“可租”到“已租”、"
        "报修从“已提交”到“维修中/已完成”）。在这种场景下，构建一个面向多角色、流程可追踪的信息系统，既能减少重复沟通，"
        "也能让关键状态变化有据可查。[8]"
    )
    h2("1.2 研究意义")
    p(
        "一方面，租赁业务天然需要“结构化数据”：房屋类型、面积、朝向楼层、租金、押金、合同金额、支付状态等都不是一句话能说清。"
        "系统化管理可以把这些要素固化为表结构与字段约束，减少口头沟通造成的信息偏差。[6][7]"
    )
    p(
        "另一方面，租赁业务也需要“过程数据”：预约时间、审核回复、维修进度、评价内容等共同构成业务闭环。"
        "当过程数据可被统计时，管理者可以更直观地发现问题，例如：哪些小区预约多但成交少、报修集中在什么类型、维修处理是否及时等，"
        "这些都为改进运营策略提供了抓手。[5][8]"
    )
    h2("1.3 国内外研究现状")
    p(
        "在信息系统实现层面，国内的租赁平台与中介管理系统大多采用 B/S 架构，后端提供统一接口，前端负责页面交互与数据展示。"
        "对毕业设计而言，研究重点不在“是否做出一个大而全的平台”，而在于把核心业务对象与流程设计清楚：房源如何发布与维护、"
        "预约如何审核与回复、合同如何落库、报修如何闭环、评价与留言如何沉淀。"
    )
    p(
        "因此，本文选择以可复现为目标：以数据库脚本和代码实现为依据，围绕真实表结构进行功能说明与理论展开；"
        "对尚未引入的高级能力（如复杂个性化推荐算法）只讨论可扩展方向，不做“已经实现”的描述。[9]"
    )
    h2("1.4 创新点与技术挑战")
    p(
        "本系统的实现重点不追求“炫技”，而是把租赁业务的关键链条做扎实：以房源为核心，串起预约、合同、报修/维修与评价，"
        "并为每个环节保留状态字段与审核回复，保证流程可追踪。"
    )
    p(
        "技术挑战主要集中在三个方面："
        "（1）多角色隔离：管理员、房主、租客看到的菜单不同，接口权限也不同；"
        "（2）状态驱动流程：例如预约、合同、报修、评价都涉及“是否审核/审核回复”，字段设计要能支撑状态流转；"
        "（3）资源管理：房源图片、报修图片等文件需要统一上传与引用，既要便于展示，也要考虑存储路径与命名策略。[1][2][5][8]"
    )
    h2("1.5 系统可行性分析")
    p(
        "技术可行性：后端采用 Spring Boot 构建 REST 接口，利用自动配置与起步依赖（starter）降低工程搭建成本；"
        "持久层使用 MyBatis-Plus 以减少重复 CRUD 代码；前端采用 Vue 进行组件化开发，并使用 Element UI 快速完成表格、表单等常见页面。"
        "这些技术栈稳定成熟，学习资料充足，适合在毕业设计周期内完成交付。[1][2][3][4]"
    )
    p(
        "数据可行性：系统使用 MySQL 存储业务数据，关系模型适合描述“房源—预约—合同”等强关联对象，"
        "并可通过主键、唯一约束与字段非空约束保证数据一致性；同时可为统计分析提供可靠的数据来源。[6][7]"
    )

    # Chapter 2
    h1("2. 技术环境概述")
    h2("2.1 开发技术")
    h2("2.1.1 编程语言")
    p(
        "后端主要采用 Java 语言实现。Java 的优势在于生态完整、工程化能力强，尤其适合“分层架构 + 明确数据模型”的业务系统。"
        "在实践中，系统通过 Controller/Service/DAO（Mapper）划分职责，使接口层、业务层与数据访问层各自保持清晰边界。[8][10]"
    )
    p(
        "前端采用 JavaScript，并以 Vue 的组件化方式组织页面。组件化并不只是“拆文件”，更重要的是把表格、表单、对话框等交互模式抽象为可复用结构，"
        "让页面随着功能扩展仍然能保持可维护性。[4]"
    )
    h2("2.1.2 数据库")
    p(
        "系统采用 MySQL 作为关系型数据库。租赁业务的核心数据（房屋信息、预约看房、合同信息、报修与维修处理、评价等）"
        "天然具有结构化与强一致性的特点，适合采用关系模型进行组织。"
        "在数据库设计中，主键用于唯一标识记录，唯一约束用于保证如“预约编号”等业务唯一性，"
        "而非空约束与默认值用于减少脏数据进入系统的概率。[6][7]"
    )
    h2("2.1.3 前后端框架")
    p(
        "后端框架方面，Spring Boot 通过自动配置机制把常见的 Web 应用能力（如 MVC、JSON 序列化、异常处理等）"
        "集成在统一的工程体系内，使开发者能把精力更多放在业务本身；"
        "而 MyBatis/MyBatis-Plus 则在“SQL 可控”和“对象映射便利”之间取得平衡，"
        "对以业务表为中心的管理系统尤其合适。[1][2][3]"
    )
    p(
        "前端框架方面，Vue 提供响应式数据与组件化组织方式；"
        "Element UI 负责将表格、分页、表单校验、对话框等管理端高频交互封装为组件；"
        "ECharts 则用于将统计结果以图表方式呈现，降低阅读成本，提升管理端决策效率。[4][5]"
    )
    h2("2.2 开发工具")
    p(
        "后端开发可使用 IntelliJ IDEA 进行代码编写、调试与依赖管理；前端可使用 VS Code 进行组件与样式开发。"
        "数据库部分可借助 Navicat 等可视化工具进行表结构查看、数据导入与调试查询，减少手工操作成本。"
    )
    h2("2.3 开发环境")
    env_parts: list[str] = []
    if runtime_versions.get("java_runtime"):
        env_parts.append(f"运行时：{runtime_versions['java_runtime']}")
    if runtime_versions.get("maven"):
        env_parts.append(f"构建工具：{runtime_versions['maven']}")
    if runtime_versions.get("node"):
        env_parts.append(f"Node.js：{runtime_versions['node']}")
    if runtime_versions.get("npm"):
        env_parts.append(f"npm：{runtime_versions['npm']}")
    if pom_versions:
        env_parts.append(
            f"后端配置：Spring Boot {pom_versions.get('spring_boot_parent','')}；java.version={pom_versions.get('java_target','')}；mysql-connector-java {pom_versions.get('mysql_connector_java','')}"
        )
    if pkg_versions:
        dep_show = []
        for k in ["vue", "element-ui", "echarts"]:
            if k in pkg_versions:
                dep_show.append(f"{k} {pkg_versions[k]}")
        if dep_show:
            env_parts.append("前端依赖：" + "，".join(dep_show))
    p("；".join(env_parts))

    # Chapter 3
    h1("3. 系统需求分析")
    h2("3.1 系统目标功能")
    h2("3.1.1 系统目标")
    p(
        "系统目标并不是简单实现“增删改查”，而是围绕租赁业务的关键节点建立一套可追溯的信息链："
        "以房源为中心，记录预约与审核、合同与支付状态、报修与维修处理、评价与反馈。"
        "当关键节点都能落库并被查询时，才算真正完成业务线上化。[8]"
    )
    h2("3.1.2 主要功能")
    p(
        "系统面向管理员、房主与租客三类角色："
        "管理员负责基础数据维护与审核监管；房主负责房源发布与租后处理；租客完成浏览、收藏、预约、合同查询与报修等操作。"
        "这种角色划分既符合业务现实，也便于后端在接口层做权限隔离。[8]"
    )
    h2("3.2 用户角色权限")
    h2("3.2.1 系统管理员")
    p(
        "管理员具备全局管理权限，主要职责是保证数据规范与流程顺畅。"
        "从功能上看，管理员需要维护房屋类型、房屋信息、用户与房主信息，审核预约看房、合同、报修、评价与“我要当房主”申请，"
        "并负责公告发布与留言管理。"
        "从管理视角看，管理员相当于“规则执行者”：通过审核字段与回复内容把控信息质量与流程节奏。"
    )
    h2("3.2.2 房主")
    p(
        "房主是房源的直接提供者，其操作紧贴业务本身：发布与维护房源信息、审核预约看房、维护合同信息、处理租后报修并填写维修进度，"
        "同时查看评价反馈以改进服务。"
        "在系统实现上，房主账号与房源、预约、合同等记录之间通过“房主账号/房主姓名”等字段关联，保证数据归属清晰。"
    )
    h2("3.2.3 租客")
    p(
        "租客最关心的是信息透明和流程可控：能快速筛选房源、提交预约并得到明确回复；"
        "签约后能随时查询合同与支付状态；入住后出现问题能提交报修并看到处理进度；"
        "服务完成后能评价并保留沟通记录。"
        "因此系统在租客侧提供收藏、预约、报修、评价与留言等入口，尽量把“需要反复问”的信息变成“随时可查”。"
    )
    h2("3.3 系统特别需求")
    p(
        "安全性方面，系统采用 Token 认证方式对接口进行访问控制：登录成功后生成 Token 并设置过期时间，"
        "后续请求在 Header 中携带 Token，由后端拦截器统一校验。"
        "在角色管理上，系统通过角色字段区分管理员/房主/租客，并在前端菜单与后端查询条件中做隔离处理。[1][2]"
    )
    p(
        "可维护性方面，系统采用分层架构与模块化设计：接口层聚焦请求参数与返回结构，业务层聚焦规则与流程，"
        "数据访问层聚焦数据库读写。这样的结构能降低修改的影响面，便于后续增加统计维度或扩展角色功能。[8]"
    )

    # Chapter 4
    h1("4. 系统总体设计")
    h2("4.1 系统架构")
    p(
        "系统采用前后端分离的 B/S 架构：前端负责页面交互与数据展示，后端负责业务逻辑处理与数据持久化，"
        "二者通过 HTTP/JSON 进行通信。"
        "这种架构的优势在于职责清晰：前端可以专注于用户体验与组件复用，后端可以专注于接口规范、权限校验与数据一致性。[8]"
    )
    p(
        "在实现层面，后端以 Controller–Service–DAO（Mapper）组织代码，"
        "对外暴露统一风格的接口返回结构；"
        "前端以页面组件承载列表、详情与表单等交互，通过 Axios 统一发起请求。"
    )
    h2("4.1.2 功能模块")
    p(
        "系统主要模块包括："
        "（1）用户与权限：管理员/房主/租客登录与会话管理；"
        "（2）基础数据：房屋类型等字典维护；"
        "（3）房屋信息：房源发布、编辑、上下架与详情展示；"
        "（4）预约看房：租客提交预约，房主/管理员审核回复；"
        "（5）合同信息：记录租期、金额、支付与审核状态；"
        "（6）报修与维修处理：报修申请、进度反馈、更新日期；"
        "（7）评价与留言：评价审核回复、留言板管理；"
        "（8）公告与轮播：公告发布与维护；"
        "（9）收藏与统计：收藏沉淀兴趣数据，统计页使用图表展示关键指标。"
    )
    h2("4.2 数据库设计")
    h2("4.2.1 表设计")
    p(
        "数据库脚本（zhihuifangwuzulin.sql）共包含 18 张表，覆盖系统配置、用户与角色、房源、预约、合同、报修与维修、评价与评论、收藏与公告等业务对象。"
        "表结构设计遵循“围绕业务对象建表、围绕流程状态设字段”的原则："
        "一张表尽量表达一个清晰的业务实体或过程记录，同时通过主键、唯一约束与默认值减少数据歧义。[6][7]"
    )
    p(
        "在关系数据库理论中，合理的表结构设计需要关注数据冗余与更新异常。"
        "因此在设计时会把相对稳定的分类信息（如房屋类型）单独成表，"
        "并通过字段引用（如房屋类型名称）与业务表关联；"
        "对于过程数据（如预约、报修），则以记录形式保留关键节点与审核字段，便于追踪与统计。[6]"
    )
    h2("4.2.2 表结构")
    p("下文给出部分关键业务表结构（字段来源于数据库脚本）。")

    table_no = 1
    for tname in schema_tables:
        cols = schema.get(tname)
        if not cols:
            continue
        meaning_map = {
            "fangwuxinxi": "房屋信息表用于存储房源的核心属性（类型、面积、状态、租金、押金等）以及图片与详情描述，是系统业务的核心数据表。",
            "yuyuekanfang": "预约看房表用于记录租客提交的预约请求与审核结果，包含预约编号、时间、租用月数等字段，便于形成可追溯的带看流程。",
            "hetongxinxi": "合同信息表用于记录签约后的关键条款与金额信息，并通过“是否支付”等字段记录支付状态，支撑合同查询与对账。",
            "fangwubaoxiu": "房屋报修表用于记录租客提交的报修申请与审核回复，结合房屋与合同信息实现租后服务入口。",
            "weixiuchuli": "维修处理表用于记录维修过程的进度、反馈与更新时间，形成报修处理闭环。",
            "woyaodangfangzhu": "我要当房主表用于记录用户提交的“成为房主”申请及审核回复，为角色扩展提供业务入口。",
        }
        p(meaning_map.get(tname, ""))
        p(f"表4.{table_no} {tname} 表结构")
        t = insert_table_before(doc, ref_para, rows=1, cols=6)
        hdr = ["字段名称", "数据类型", "字段长度", "是否为空", "默认值", "说明"]
        for i, h in enumerate(hdr):
            t.rows[0].cells[i].text = h
        for col in cols:
            row = t.add_row().cells
            row[0].text = col["name"]
            row[1].text = col["type"]
            row[2].text = col["length"]
            row[3].text = col["nullable"]
            row[4].text = col["default"]
            row[5].text = col["comment"]
        table_no += 1

    # Chapter 5
    h1("5. 系统详细实现")
    h2("5.1 登录认证与 Token 机制")
    p(
        "系统的登录认证采用“Token + 拦截器”的方式实现。登录成功后，后端生成随机 Token 并写入 token 表，"
        "其中包含用户标识、角色、关联表与过期时间；前端保存该 Token，并在后续请求的 Header 中携带。"
        "后端拦截器统一读取 Header 的 Token，校验有效性后把用户信息写入 Session，用于后续的权限判断与数据隔离。[1][2]"
    )
    p(
        "这种设计的好处是把认证逻辑集中在拦截器层处理，业务接口可以更专注于“做什么”；"
        "同时 Token 设置过期时间，可以降低长期有效会话带来的风险。"
        "需要说明的是，当前系统密码校验采用明文比对，属于可改进点：在实际部署时应引入加盐哈希等方式保护密码安全。[2][8]"
    )
    h2("5.2 文件上传与图片管理")
    p(
        "房源图片与报修图片是租赁系统的高频资源。系统提供统一文件上传接口接收 MultipartFile，"
        "将文件保存至后端静态资源目录的 upload 文件夹，并返回文件名供业务表保存引用。"
        "在展示时，前端根据保存的图片路径进行渲染，保证“数据记录”和“文件资源”之间有稳定关联。[1][2]"
    )
    p(
        "文件上传属于典型的“方便与安全并存”的功能：为了便于管理，系统采用时间戳命名减少重名冲突；"
        "为了更安全的工程落地，还可以在此基础上增加文件类型白名单、大小限制与访问权限控制等策略。[2][8]"
    )
    h2("5.3 房屋信息与预约流程")
    p(
        "房屋信息模块是系统的入口：房主/管理员维护房源的类型、面积、朝向楼层、状态、租金与押金等关键属性，"
        "并上传图片与详情描述。"
        "这些字段并非“越多越好”，而是围绕租客决策所需的信息组织：能否租、多少钱、位置与配套、房屋设施是否满足需求。"
    )
    p(
        "预约看房模块对应真实业务中的“带看”环节。租客在浏览房源详情后提交预约，"
        "预约记录包含预约编号、时间、租用月数等信息，并通过审核状态与审核回复字段记录结果。"
        "当预约信息可被查询与统计时，带看安排就不会只停留在聊天记录里，后续也能复盘预约转化情况。[6][8]"
    )
    h2("5.4 合同、报修与评价闭环")
    p(
        "合同信息用于固化租赁关系。系统通过合同表记录合同编号、租用月数、租用金额、押金、合同金额等关键条款，"
        "并用“是否支付/审核状态”等字段支撑管理端对合同状态的判断。"
        "在数据层面，这些字段为后续统计（如租金汇总）提供了直接的数据来源。[6][7]"
    )
    p(
        "租后服务以报修为入口：租客提交报修内容与图片，房主/管理员可审核回复；"
        "维修处理模块则进一步记录维修进度与反馈，直到状态完成。"
        "评价模块把租客体验沉淀为可追踪信息，形成从“问题—处理—反馈”的闭环。"
    )
    h2("5.5 统计与可视化")
    p(
        "当业务数据落库后，可视化的价值才真正体现出来。后台统计页面对租赁、报修等数据进行汇总，"
        "并使用 ECharts 绘制柱状图、折线图或饼图等常见图表，"
        "让管理者能在较短时间内把握整体情况与趋势变化。[5]"
    )
    p(
        "值得强调的是，可视化并不等于“图越炫越好”。更重要的是选取贴合业务的指标，例如："
        "不同房屋类型的发布数量、预约数量变化、报修类型分布等。"
        "图表的作用是把“表格里不容易看出来的规律”呈现出来，为后续优化提供方向。[5][8]"
    )

    # Chapter 6
    h1("6. 系统测试")
    h2("6.1 测试目的")
    p(
        "系统测试的目的不仅是“能跑起来”，更重要的是验证核心流程在常见输入与异常输入下都能得到可预期的结果。"
        "对租赁管理系统而言，测试应覆盖登录鉴权、房源维护、预约审核、合同维护、报修与维修处理等关键路径，"
        "同时关注数据一致性（例如同一预约编号的唯一性、审核状态字段的合理取值等）。[8]"
    )
    h2("6.2 测试数据与用例设计")
    p(
        "数据库脚本包含初始化数据，可用于复现实验与功能验证。"
        "在测试用例设计中，可以基于“等价类划分、边界值分析”等方法覆盖典型场景："
        "例如租用月数为 0/负数/超大值时的校验，必填字段为空时的提示等。[8]"
    )
    p("关键表插入记录数统计如下（由数据库脚本统计得到）：")
    stats = insert_table_before(doc, ref_para, rows=1, cols=3)
    stats.rows[0].cells[0].text = "表名"
    stats.rows[0].cells[1].text = "含义"
    stats.rows[0].cells[2].text = "插入记录数（脚本统计）"
    for tname, meaning in key_tables:
        row = stats.add_row().cells
        row[0].text = tname
        row[1].text = meaning
        row[2].text = str(insert_counts.get(tname, 0))
    h2("6.2.1 功能测试")
    p(
        "建议按如下用例验证核心流程："
        "（1）管理员/房主/租客登录并获取 Token；"
        "（2）房主发布房源，租客浏览并收藏；"
        "（3）租客提交预约看房，房主或管理员审核并回复；"
        "（4）生成/维护合同记录并查询支付状态；"
        "（5）租客提交报修，房主填写维修处理记录；"
        "（6）租客提交评价与留言，管理员维护公告。"
    )
    h2("6.2.2 兼容性测试")
    p("前端建议在 Chrome、Edge 等主流浏览器下进行页面兼容性验证，并关注不同分辨率下的表格与表单布局。")
    h2("6.3 测试结果分析")
    p(
        "在本机环境中，后端已完成 Maven 构建（mvn -DskipTests package），能够生成后端可运行包；"
        "数据库脚本可用于初始化并支撑业务流程验证。前端可通过 npm scripts 进行本地启动与构建。"
        "在后续优化中，可继续补充自动化测试与更细粒度的数据校验，提升系统在异常输入下的鲁棒性。[8]"
    )

    # Chapter 7
    h1("7. 结论")
    p(
        "本文围绕智慧房屋租赁业务流程完成了系统需求分析、总体设计、数据库设计与关键功能实现，形成面向管理员、房主与租客的多角色管理系统。"
        "后续可在更完善的运行环境下进一步开展自动化测试与性能压测，并在现有收藏/行为数据基础上引入更精细的个性化推荐策略。[8][9]"
    )

    # Replace references with domestic books
    idx_ref = find_para_index_exact(doc, "参考文献")
    idx_ack = find_para_index_exact(doc, "致谢")
    if idx_ref != -1 and idx_ack != -1 and idx_ack > idx_ref:
        for i in range(idx_ack - 1, idx_ref, -1):
            delete_paragraph(doc.paragraphs[i])
        ack_para = doc.paragraphs[find_para_index_exact(doc, "致谢")]
        refs = [
            "[1] 汪云飞. JavaEE开发的颠覆者：Spring Boot实战[M]. 北京: 电子工业出版社, 2016.",
            "[2] 黑马程序员. Spring Boot企业级开发教程[M]. 北京: 人民邮电出版社, 2019.",
            "[3] 赖帆. MyBatis核心技术全解与项目实战[M]. 北京: 人民邮电出版社, 2024.",
            "[4] 霍春阳. Vue.js设计与实现[M]. 北京: 人民邮电出版社, 2022.",
            "[5] 范路桥, 张良均, 郑述招, 肖秀娟, 李明. Web数据可视化（ECharts版）[M]. 北京: 人民邮电出版社, 2021.",
            "[6] 王珊, 萨师煊. 数据库系统概论（第5版）[M]. 北京: 高等教育出版社, 2014.",
            "[7] Forta B. MySQL必知必会（第2版）[M]. 北京: 人民邮电出版社, 2024.",
            "[8] 张海藩, 牟永敏. 软件工程导论（第6版）[M]. 北京: 清华大学出版社, 2013.",
            "[9] 项亮. 推荐系统实践[M]. 北京: 人民邮电出版社, 2012.",
            "[10] 周志明. 深入理解Java虚拟机：JVM高级特性与最佳实践（第3版）[M]. 北京: 机械工业出版社, 2019.",
        ]
        for line in reversed(refs):
            p_ref = ack_para.insert_paragraph_before(line)
            p_ref.style = "List Paragraph"

    # Replace acknowledgement content (avoid fake names)
    idx_ack = find_para_index_exact(doc, "致谢")
    if idx_ack != -1 and idx_ack + 1 < len(doc.paragraphs):
        doc.paragraphs[idx_ack + 1].text = (
            "在本毕业设计的完成过程中，感谢指导老师在选题、技术路线与论文写作方面给予的指导与建议；"
            "感谢同学与朋友在系统调试与资料整理过程中提供的帮助。"
        )

    doc.save(str(out_path))
    doc.save(str(work_copy))
    print(f"Saved: {out_path}")
    print(f"Saved: {work_copy}")


if __name__ == "__main__":
    main()
