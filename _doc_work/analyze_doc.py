from docx import Document
from pathlib import Path

THESIS = Path(r"D:\BaiduNetdiskDownload\T020源码\T020源码\_doc_work\thesis.docx")
SPEC = Path(r"D:\BaiduNetdiskDownload\T020源码\T020源码\_doc_work\functional_spec.docx")


def iter_blocks(doc):
    # yields (kind, obj)
    for p in doc.paragraphs:
        yield ("p", p)
    for t in doc.tables:
        yield ("t", t)


def get_heading_level(style_name: str):
    if not style_name:
        return None
    if style_name.startswith("Heading"):
        try:
            return int(style_name.split()[-1])
        except Exception:
            return None
    if style_name.startswith("标题"):
        # common: 标题 1, 标题 2 ...
        parts = style_name.split()
        if parts and parts[-1].isdigit():
            return int(parts[-1])
        return None
    return None


def section_stats(doc):
    sections=[]
    current=None
    for p in doc.paragraphs:
        text=(p.text or '').strip()
        if not text:
            continue
        lvl=get_heading_level(getattr(p.style,'name', ''))
        if lvl is not None:
            current={"level":lvl,"title":text,"paras":0,"chars":0}
            sections.append(current)
            continue
        if current is not None:
            current["paras"] += 1
            current["chars"] += len(text)
    return sections


def dump_top(sections, n=40):
    for s in sections[:n]:
        print(f"L{s['level']} {s['title']} | paras={s['paras']} chars={s['chars']}")


def main():
    thesis=Document(str(THESIS))
    spec=Document(str(SPEC))
    print("THESIS sections:")
    dump_top(section_stats(thesis), 60)
    print("\nSPEC sections:")
    dump_top(section_stats(spec), 60)

if __name__=='__main__':
    main()
