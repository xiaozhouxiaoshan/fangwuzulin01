from docx import Document
from pathlib import Path

spec_path = Path(r'D:\BaiduNetdiskDownload\T020源码\T020源码\_doc_work\functional_spec.docx')
doc = Document(str(spec_path))

def heading_level(style_name: str):
    if not style_name:
        return None
    if style_name.startswith('Heading'):
        try:
            return int(style_name.split()[-1])
        except Exception:
            return None
    if style_name.startswith('标题'):
        parts = style_name.split()
        if parts and parts[-1].isdigit():
            return int(parts[-1])
    return None


def extract(title, max_lines=200):
    start=None
    for i,p in enumerate(doc.paragraphs):
        if (p.text or '').strip()==title and heading_level(p.style.name if p.style else '') is not None:
            start=i
            break
    if start is None:
        print('NOT FOUND', title)
        return
    lvl=heading_level(doc.paragraphs[start].style.name)
    print('===', title, '===')
    lines=0
    for p in doc.paragraphs[start+1:]:
        t=(p.text or '').strip()
        if not t:
            continue
        plvl=heading_level(p.style.name if p.style else '')
        if plvl is not None and plvl<=lvl:
            break
        print(t)
        lines += 1
        if lines>=max_lines:
            break
    print()

for t in ['5.1管理员功能模块','5.2房主功能模块','5.3用户功能模块','3.3系统功能设计']:
    extract(t)
