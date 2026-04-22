from docx import Document
from pathlib import Path

path = Path(r'D:\BaiduNetdiskDownload\T020源码\T020源码\_doc_work\thesis.docx')
doc = Document(str(path))

def heading_level(p):
    if not p.style:
        return None
    name = p.style.name
    if name.startswith('Heading'):
        try:
            return int(name.split()[-1])
        except Exception:
            return None
    return None


def extract(title, max_paras=15):
    start=None
    for i,p in enumerate(doc.paragraphs):
        if (p.text or '').strip() == title and heading_level(p) is not None:
            start=i
            break
    if start is None:
        print('NOT FOUND', title)
        return
    lvl = heading_level(doc.paragraphs[start])
    print('\n===', title, '===')
    out=[]
    for p in doc.paragraphs[start+1:]:
        t=(p.text or '').strip()
        if not t:
            continue
        plvl = heading_level(p)
        if plvl is not None and plvl <= lvl:
            break
        out.append(t)
        if len(out) >= max_paras:
            break
    for t in out:
        print('-', t)

for t in [
    'Spring Boot 框架',
    'MyBatis / MyBatis-Plus 持久层技术',
    '系统总体架构设计',
    '3.2.1 管理员功能需求',
    '3.2.2 租客功能需求',
    '系统功能模块设计',
    '4.3.3 数据表设计',
    'Token 拦截器设计',
    '文件上传设计',
    '智慧推荐模型设计',
    '留言与论坛模块设计与实现',
    '参考文献'
]:
    extract(t)
